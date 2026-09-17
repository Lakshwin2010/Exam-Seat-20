import json
import base64
import datetime
from io import BytesIO
from pathlib import Path
from typing import Dict, List, Optional, Any
from fastapi import APIRouter, Depends, HTTPException, Query, Response
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
import urllib.parse

from app.database import get_db
from app.config import DATA_DIR
from app import models

router = APIRouter(prefix="/api/email", tags=["email"])

RECIPIENTS_FILE = DATA_DIR / "email_recipients.json"

DEFAULT_RECIPIENTS = {
    # Grade XI Sections (8 sections)
    "XI - A": {"email": "erd_s1401728@csacademy.in", "teacherName": "Class Teacher (XI-A)"},
    "XI - B": {"email": "erd_s1803155@csacademy.in", "teacherName": "Class Teacher (XI-B)"},
    "XI - C": {"email": "", "teacherName": "Class Teacher (XI-C)"},
    "XI - D": {"email": "", "teacherName": "Class Teacher (XI-D)"},
    "XI - E": {"email": "", "teacherName": "Class Teacher (XI-E)"},
    "XI - F": {"email": "", "teacherName": "Class Teacher (XI-F)"},
    "XI - G": {"email": "", "teacherName": "Class Teacher (XI-G)"},
    "XI - H": {"email": "", "teacherName": "Class Teacher (XI-H)"},
    # Grade XII Sections (7 sections)
    "XII - A": {"email": "", "teacherName": "Class Teacher (XII-A)"},
    "XII - B": {"email": "", "teacherName": "Class Teacher (XII-B)"},
    "XII - C": {"email": "", "teacherName": "Class Teacher (XII-C)"},
    "XII - D": {"email": "", "teacherName": "Class Teacher (XII-D)"},
    "XII - E": {"email": "", "teacherName": "Class Teacher (XII-E)"},
    "XII - F": {"email": "", "teacherName": "Class Teacher (XII-F)"},
    "XII - G": {"email": "", "teacherName": "Class Teacher (XII-G)"},
}

def load_recipients() -> Dict[str, Dict[str, str]]:
    if not RECIPIENTS_FILE.exists():
        save_recipients(DEFAULT_RECIPIENTS)
        return DEFAULT_RECIPIENTS
    try:
        with open(RECIPIENTS_FILE, "r", encoding="utf-8") as f:
            data = json.load(f)
            # Ensure defaults exist
            for k, v in DEFAULT_RECIPIENTS.items():
                if k not in data:
                    data[k] = v
            return data
    except Exception:
        return DEFAULT_RECIPIENTS

def save_recipients(data: Dict[str, Dict[str, str]]):
    with open(RECIPIENTS_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2)

@router.get("/recipients")
def get_email_recipients(
    session_id: Optional[str] = Query(None, description="Optional session ID to compute live seating counts"),
    db: Session = Depends(get_db)
):
    """Returns recipient configurations for all sections with student count and live allocation counts."""
    recipients_map = load_recipients()
    
    session = None
    if session_id:
        session = db.query(models.ExamSession).filter(models.ExamSession.id == session_id).first()
    if not session:
        alloc = db.query(models.SeatAllocation).first()
        if alloc:
            session = alloc.session
        else:
            session = db.query(models.ExamSession).first()

    all_sections = [
        "XI - A", "XI - B", "XI - C", "XI - D", "XI - E", "XI - F", "XI - G", "XI - H",
        "XII - A", "XII - B", "XII - C", "XII - D", "XII - E", "XII - F", "XII - G"
    ]
    result = []
    
    for sec in all_sections:
        student_count = db.query(models.Student).filter(models.Student.grade == sec).count()
        cfg = recipients_map.get(sec, {"email": "", "teacherName": f"Class Teacher ({sec})"})

        room_seats_count = 0
        xi_seats_count = 0
        xii_seats_count = 0
        despatch_count = 0

        if session:
            room = db.query(models.ExamRoom).filter(models.ExamRoom.name == sec).first()
            if not room:
                room = db.query(models.ExamRoom).filter(models.ExamRoom.name.ilike(f"%{sec}%")).first()
            
            if room:
                room_allocs = db.query(models.SeatAllocation).filter(
                    models.SeatAllocation.session_id == session.id,
                    models.SeatAllocation.room_id == room.id
                ).all()
                room_seats_count = len(room_allocs)
                for ra in room_allocs:
                    if ra.student:
                        ug = ra.student.grade.upper()
                        if "XI" in ug and "XII" not in ug:
                            xi_seats_count += 1
                        elif "XII" in ug:
                            xii_seats_count += 1

            class_students_ids = [s.id for s in db.query(models.Student.id).filter(models.Student.grade == sec).all()]
            if class_students_ids:
                despatch_count = db.query(models.SeatAllocation).filter(
                    models.SeatAllocation.session_id == session.id,
                    models.SeatAllocation.student_id.in_(class_students_ids)
                ).count()

        result.append({
            "section": sec,
            "gradeLevel": "Grade XI" if ("XI" in sec and "XII" not in sec) else "Grade XII",
            "studentCount": student_count,
            "roomSeatsCount": room_seats_count,
            "xiSeatsCount": xi_seats_count,
            "xiiSeatsCount": xii_seats_count,
            "despatchCount": despatch_count,
            "email": cfg.get("email", ""),
            "teacherName": cfg.get("teacherName", f"Class Teacher ({sec})"),
            "isConfigured": bool(cfg.get("email", "").strip())
        })
        
    return {
        "status": "success",
        "sessionId": session.id if session else None,
        "sessionName": session.name if session else None,
        "totalSections": len(result),
        "sections": result
    }

@router.post("/recipients")
def update_email_recipients(payload: Dict[str, Any]):
    """Updates recipient emails for one or more sections."""
    current = load_recipients()
    updates = payload.get("recipients", {})
    for sec, info in updates.items():
        if isinstance(info, str):
            if sec not in current:
                current[sec] = {"teacherName": f"Class Teacher ({sec})", "email": ""}
            current[sec]["email"] = info.strip()
        elif isinstance(info, dict):
            current[sec] = {
                "email": info.get("email", "").strip(),
                "teacherName": info.get("teacherName", f"Class Teacher ({sec})")
            }
    save_recipients(current)
    return {"status": "success", "recipients": current}

def generate_class_excel(grade_class: str, session_id: Optional[str], db: Session) -> tuple:
    # Normalize grade class string (e.g. 'XI-A' -> 'XI - A')
    norm = grade_class.strip()
    if "-" in norm and " - " not in norm:
        parts = norm.split("-")
        norm = f"{parts[0].strip()} - {parts[1].strip()}"
        
    sec_slug = norm.replace(" ", "")

    session = None
    if session_id:
        session = db.query(models.ExamSession).filter(models.ExamSession.id == session_id).first()
    if not session:
        alloc = db.query(models.SeatAllocation).first()
        if alloc:
            session = alloc.session
        else:
            session = db.query(models.ExamSession).first()

    session_name = session.name if session else "Exam Session"
    session_time = f"{session.date} ({session.time_slot})" if session else datetime.datetime.now().strftime("%d-%b-%Y")

    # 1. Gather Room Seating (Who will be in this classroom: both 11th & 12th)
    room = db.query(models.ExamRoom).filter(models.ExamRoom.name == norm).first()
    if not room:
        room = db.query(models.ExamRoom).filter(models.ExamRoom.name.ilike(f"%{norm}%")).first()

    room_allocations = []
    if room and session:
        room_allocations = (
            db.query(models.SeatAllocation)
            .filter(
                models.SeatAllocation.room_id == room.id,
                models.SeatAllocation.session_id == session.id
            )
            .order_by(
                models.SeatAllocation.row.asc(),
                models.SeatAllocation.col.asc(),
                models.SeatAllocation.seat_label.asc()
            )
            .all()
        )

    # 2. Gather Home Class Students (Where that class students will go)
    class_students = (
        db.query(models.Student)
        .filter(models.Student.grade == norm)
        .order_by(models.Student.roll_no.asc())
        .all()
    )
    if not class_students:
        class_students = (
            db.query(models.Student)
            .filter(models.Student.grade.ilike(f"%{grade_class.strip()}%"))
            .order_by(models.Student.roll_no.asc())
            .all()
        )

    stud_ids = [s.id for s in class_students]
    alloc_by_stud_id = {}
    if session and stud_ids:
        allocs = db.query(models.SeatAllocation).filter(
            models.SeatAllocation.session_id == session.id,
            models.SeatAllocation.student_id.in_(stud_ids)
        ).all()
        for a in allocs:
            alloc_by_stud_id[a.student_id] = a

    wb = openpyxl.Workbook()
    
    # Common Styling Definitions
    title_font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="475569")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill_blue = PatternFill(start_color="1D4ED8", end_color="1D4ED8", fill_type="solid")
    header_fill_emerald = PatternFill(start_color="047857", end_color="047857", fill_type="solid")
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    regular_font = Font(name="Calibri", size=10)
    bold_font = Font(name="Calibri", size=10, bold=True)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )

    # ═════════════════════════════════════════════════════════════════
    # TAB 1: ROOM IN-HALL SEATING PLAN ("Who will be in that class")
    # ═════════════════════════════════════════════════════════════════
    ws_room = wb.active
    ws_room.title = f"Room_{sec_slug}_Seating"[:31]
    ws_room.views.sheetView[0].showGridLines = True

    xi_count = sum(1 for a in room_allocations if a.student and ("XI" in a.student.grade.upper() and "XII" not in a.student.grade.upper()))
    xii_count = sum(1 for a in room_allocations if a.student and ("XII" in a.student.grade.upper()))

    ws_room.merge_cells("A1:G1")
    ws_room["A1"] = f"CS ACADEMY - IN-ROOM EXAM SEATING PLAN (ROOM {norm})"
    ws_room["A1"].font = title_font
    ws_room["A1"].alignment = center_align

    ws_room.merge_cells("A2:G2")
    ws_room["A2"] = (
        f"Session: {session_name} | {session_time} | "
        f"Total Seated: {len(room_allocations)} (Class XI: {xi_count}, Class XII: {xii_count}) | "
        f"Room Capacity: {room.capacity if room else 30}"
    )
    ws_room["A2"].font = subtitle_font
    ws_room["A2"].alignment = center_align

    ws_room.row_dimensions[1].height = 24
    ws_room.row_dimensions[2].height = 18
    ws_room.row_dimensions[4].height = 24

    room_headers = ["Desk / Seat", "Exam / Roll No", "Student Name", "Class / Section", "Subject Code", "Subject Name", "Special Needs"]
    for col_num, h in enumerate(room_headers, 1):
        cell = ws_room.cell(row=4, column=col_num, value=h)
        cell.font = header_font
        cell.fill = header_fill_blue
        cell.alignment = center_align
        cell.border = thin_border

    if room_allocations:
        for idx, alloc in enumerate(room_allocations, 1):
            curr_row = 4 + idx
            ws_room.row_dimensions[curr_row].height = 20
            st = alloc.student
            sub = alloc.subject

            row_vals = [
                alloc.seat_label,
                st.roll_no if st else "Vacant",
                st.name if st else "-",
                st.grade if st else "-",
                sub.code if sub else "-",
                sub.name if sub else "-",
                "Yes (Priority Front)" if (st and st.special_needs) else "No"
            ]

            is_alt = (idx % 2 == 0)
            for c_idx, val in enumerate(row_vals, 1):
                c = ws_room.cell(row=curr_row, column=c_idx, value=val)
                c.font = bold_font if c_idx in (1, 2, 3) else regular_font
                c.border = thin_border
                if is_alt:
                    c.fill = alt_fill
                c.alignment = center_align if c_idx in (1, 2, 4, 5, 7) else left_align
    else:
        ws_room.merge_cells("A5:G5")
        ws_room["A5"] = "No seating allocations generated yet for this room and session. Run seating allocation in the ExamHall portal."
        ws_room["A5"].alignment = center_align
        ws_room["A5"].font = subtitle_font

    for col in ws_room.columns:
        max_l = 0
        col_let = get_column_letter(col[0].column)
        for c in col:
            if c.row > 2 and c.value:
                max_l = max(max_l, len(str(c.value)))
        ws_room.column_dimensions[col_let].width = max(max_l + 4, 12)

    # ═════════════════════════════════════════════════════════════════
    # TAB 2: CLASS OUTWARD DESPATCH ("Where that class students will go")
    # ═════════════════════════════════════════════════════════════════
    ws_despatch = wb.create_sheet(title=f"Class_{sec_slug}_Despatch"[:31])
    ws_despatch.views.sheetView[0].showGridLines = True

    ws_despatch.merge_cells("A1:G1")
    ws_despatch["A1"] = f"CS ACADEMY - CLASS OUTWARD EXAM DESPATCH (CLASS {norm})"
    ws_despatch["A1"].font = title_font
    ws_despatch["A1"].alignment = center_align

    ws_despatch.merge_cells("A2:G2")
    ws_despatch["A2"] = (
        f"Home Class: {norm} | Session: {session_name} ({session_time}) | "
        f"Total Students: {len(class_students)} | Exported: {datetime.datetime.now().strftime('%d-%b-%Y %I:%M %p')}"
    )
    ws_despatch["A2"].font = subtitle_font
    ws_despatch["A2"].alignment = center_align

    ws_despatch.row_dimensions[1].height = 24
    ws_despatch.row_dimensions[2].height = 18
    ws_despatch.row_dimensions[4].height = 24

    despatch_headers = ["S.No.", "Roll / Exam No", "Student Name", "Gender", "Exam Subject", "Assigned Exam Hall", "Desk / Seat"]
    for col_num, h in enumerate(despatch_headers, 1):
        cell = ws_despatch.cell(row=4, column=col_num, value=h)
        cell.font = header_font
        cell.fill = header_fill_emerald
        cell.alignment = center_align
        cell.border = thin_border

    for idx, st in enumerate(class_students, 1):
        curr_row = 4 + idx
        ws_despatch.row_dimensions[curr_row].height = 20
        st_alloc = alloc_by_stud_id.get(st.id)

        if st_alloc and st_alloc.subject:
            subject_str = f"{st_alloc.subject.code} - {st_alloc.subject.name}"
        elif st.enrolled_subjects:
            subject_str = ", ".join([s.code for s in st.enrolled_subjects])
        else:
            subject_str = "-"

        room_str = st_alloc.room.name if (st_alloc and st_alloc.room) else "Not Allocated"
        seat_str = st_alloc.seat_label if st_alloc else "-"

        row_vals = [
            idx,
            st.roll_no,
            st.name,
            st.gender or "M",
            subject_str,
            room_str,
            seat_str
        ]

        is_alt = (idx % 2 == 0)
        for c_idx, val in enumerate(row_vals, 1):
            c = ws_despatch.cell(row=curr_row, column=c_idx, value=val)
            c.font = bold_font if c_idx in (2, 3, 6, 7) else regular_font
            c.border = thin_border
            if is_alt:
                c.fill = alt_fill
            c.alignment = center_align if c_idx in (1, 2, 4, 6, 7) else left_align

    for col in ws_despatch.columns:
        max_l = 0
        col_let = get_column_letter(col[0].column)
        for c in col:
            if c.row > 2 and c.value:
                max_l = max(max_l, len(str(c.value)))
        ws_despatch.column_dimensions[col_let].width = max(max_l + 4, 12)

    filename = f"Exam_Seating_and_Despatch_{sec_slug}.xlsx"
    return wb, filename, len(class_students), len(room_allocations), xi_count, xii_count

@router.get("/class-sheet/{grade_class}")
def get_class_sheet(
    grade_class: str,
    session_id: Optional[str] = Query(None, description="Optional exam session ID"),
    download: bool = Query(False, description="If true, returns binary excel file download"),
    db: Session = Depends(get_db)
):
    """Generates and returns the clean 2-tab Excel workbook for a class section (as base64 or direct file download)."""
    wb, filename, student_count, room_seats_count, xi_seats_count, xii_seats_count = generate_class_excel(grade_class, session_id, db)
    
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    xlsx_bytes = bio.getvalue()

    if download:
        return Response(
            content=xlsx_bytes,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'}
        )

    base64_data = base64.b64encode(xlsx_bytes).decode("utf-8")
    
    # Also generate the companion Dispatch PDF so callers have both formats
    try:
        pdf_bytes, pdf_filename, _, _ = generate_class_dispatch_pdf(grade_class, session_id, db)
        pdf_base64 = base64.b64encode(pdf_bytes).decode("utf-8")
    except Exception as e:
        pdf_base64 = None
        pdf_filename = f"Exam_Dispatch_{grade_class.replace(' ', '')}.pdf"

    return {
        "status": "success",
        "section": grade_class,
        "filename": filename,
        "studentCount": student_count,
        "despatchCount": student_count,
        "roomSeatsCount": room_seats_count,
        "xiSeatsCount": xi_seats_count,
        "xiiSeatsCount": xii_seats_count,
        "base64": base64_data,
        "pdfBase64": pdf_base64,
        "pdfFilename": pdf_filename,
        "sizeBytes": len(xlsx_bytes)
    }

def generate_class_dispatch_pdf(grade_class: str, session_id: Optional[str], db: Session) -> tuple:
    """Generates the official vector-clean Class Outward Exam Dispatch PDF using ReportLab."""
    norm = grade_class.strip()
    if "-" in norm and " - " not in norm:
        parts = norm.split("-")
        norm = f"{parts[0].strip()} - {parts[1].strip()}"
    sec_slug = norm.replace(" ", "")

    session = None
    if session_id:
        session = db.query(models.ExamSession).filter(models.ExamSession.id == session_id).first()
    if not session:
        alloc = db.query(models.SeatAllocation).first()
        if alloc:
            session = alloc.session
        else:
            session = db.query(models.ExamSession).first()

    session_name = session.name if session else "Exam Session"
    session_date = session.date if session else datetime.datetime.now().strftime("%d-%b-%Y")
    time_slot_str = session.time_slot if (session and session.time_slot) else "Morning (09:30 AM - 12:30 PM)"

    # Home Class Students
    class_students = (
        db.query(models.Student)
        .filter(models.Student.grade == norm)
        .order_by(models.Student.roll_no.asc())
        .all()
    )
    if not class_students:
        class_students = (
            db.query(models.Student)
            .filter(models.Student.grade.ilike(f"%{grade_class.strip()}%"))
            .order_by(models.Student.roll_no.asc())
            .all()
        )

    stud_ids = [s.id for s in class_students]
    alloc_by_stud_id = {}
    if session and stud_ids:
        allocs = db.query(models.SeatAllocation).filter(
            models.SeatAllocation.session_id == session.id,
            models.SeatAllocation.student_id.in_(stud_ids)
        ).all()
        for a in allocs:
            alloc_by_stud_id[a.student_id] = a

    # Calculate room breakdown
    room_splits = {}
    subjects_set = set()
    for st in class_students:
        al = alloc_by_stud_id.get(st.id)
        r_name = al.room.name if (al and al.room) else "Unassigned"
        if r_name not in room_splits:
            room_splits[r_name] = {"room": r_name, "count": 0, "rolls": [], "names": []}
        room_splits[r_name]["count"] += 1
        room_splits[r_name]["rolls"].append(st.roll_no)
        room_splits[r_name]["names"].append(st.name)
        if al and al.subject:
            subjects_set.add(al.subject.name)

    room_summary_items = []
    for r_name, info in room_splits.items():
        rolls = [int(r) for r in info["rolls"] if str(r).isdigit()]
        if rolls:
            roll_str = f"Roll {min(rolls)} to {max(rolls)}"
        else:
            roll_str = f"{info['count']} students"
        room_summary_items.append(f"{r_name}: {info['count']} Students ({roll_str})")

    room_summary_text = "   |   ".join(room_summary_items) if room_summary_items else "No room assignments"
    subject_summary = ", ".join(sorted(subjects_set)) if subjects_set else "Standard Curriculum"

    # Build PDF with ReportLab
    bio = BytesIO()
    doc = SimpleDocTemplate(
        bio,
        pagesize=A4,
        leftMargin=24,
        rightMargin=24,
        topMargin=20,
        bottomMargin=20
    )

    styles = getSampleStyleSheet()

    school_style = ParagraphStyle(
        'SchoolHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        textColor=colors.HexColor('#1E3A8A'),
        alignment=TA_CENTER
    )
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10.5,
        leading=13,
        textColor=colors.HexColor('#0F172A'),
        alignment=TA_CENTER
    )
    meta_val_style = ParagraphStyle(
        'MetaVal',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=10.5,
        textColor=colors.HexColor('#0F172A')
    )
    banner_style = ParagraphStyle(
        'BannerText',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor('#1E3A8A'),
        alignment=TA_CENTER
    )
    th_style = ParagraphStyle(
        'TH',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.white,
        alignment=TA_CENTER
    )
    cell_style = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor('#1E293B')
    )
    cell_bold_style = ParagraphStyle(
        'TableCellBold',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor('#1E293B')
    )
    cell_room_style = ParagraphStyle(
        'TableCellRoom',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#1D4ED8')
    )
    instruction_style = ParagraphStyle(
        'InstText',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7,
        leading=9,
        textColor=colors.HexColor('#64748B')
    )

    story = []
    story.append(Paragraph('CS ACADEMY HIGHER SECONDARY SCHOOL', school_style))
    story.append(Spacer(1, 2))
    story.append(Paragraph(f'OFFICIAL CLASS OUTWARD EXAM DESPATCH & SEATING LIST — {norm}', title_style))
    story.append(Spacer(1, 5))

    # Meta grid
    meta_data = [
        [
            Paragraph(f'<b>Session:</b> {session_name}', meta_val_style),
            Paragraph(f'<b>Date:</b> {session_date}', meta_val_style),
            Paragraph(f'<b>Class:</b> {norm} ({len(class_students)} Students)', meta_val_style),
            Paragraph(f'<b>Time:</b> {time_slot_str}', meta_val_style)
        ]
    ]
    meta_table = Table(meta_data, colWidths=[140, 125, 142, 140])
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F1F5F9')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#CBD5E1')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
        ('TOPPADDING', (0,0), (-1,-1), 3.5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3.5),
        ('LEFTPADDING', (0,0), (-1,-1), 6),
        ('RIGHTPADDING', (0,0), (-1,-1), 6),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 4))

    # Room summary alert bar
    summary_data = [[
        Paragraph(f'<b>📍 Room Allocation Split:</b> {room_summary_text}', banner_style)
    ]]
    summary_table = Table(summary_data, colWidths=[547])
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#EFF6FF')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#BFDBFE')),
        ('TOPPADDING', (0,0), (-1,-1), 3),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3),
        ('LEFTPADDING', (0,0), (-1,-1), 6),
        ('RIGHTPADDING', (0,0), (-1,-1), 6),
        ('ALIGN', (0,0), (-1,-1), 'CENTER')
    ]))
    story.append(summary_table)
    story.append(Spacer(1, 6))

    # Student Table
    col_widths = [26, 54, 150, 32, 110, 110, 65]
    table_data = [[
        Paragraph('S.No', th_style),
        Paragraph('Roll No', th_style),
        Paragraph('Student Name', th_style),
        Paragraph('Gender', th_style),
        Paragraph('Exam Subject', th_style),
        Paragraph('Assigned Exam Hall', th_style),
        Paragraph('Desk / Seat', th_style)
    ]]

    for idx, st in enumerate(class_students, 1):
        al = alloc_by_stud_id.get(st.id)
        room_txt = al.room.name if (al and al.room) else 'Unassigned'
        seat_txt = al.seat_label if al else '-'
        if al and al.subject:
            sub_txt = f"{al.subject.code} - {al.subject.name}"
        elif st.enrolled_subjects:
            sub_txt = ", ".join([s.code for s in st.enrolled_subjects])
        else:
            sub_txt = "-"

        table_data.append([
            Paragraph(str(idx), cell_style),
            Paragraph(str(st.roll_no), cell_bold_style),
            Paragraph(st.name, cell_bold_style),
            Paragraph(st.gender or 'M', cell_style),
            Paragraph(sub_txt, cell_style),
            Paragraph(room_txt, cell_room_style),
            Paragraph(seat_txt, cell_bold_style)
        ])

    t = Table(table_data, colWidths=col_widths, repeatRows=1)
    ts = [
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1D4ED8')),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
        ('TOPPADDING', (0,0), (-1,-1), 2.5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 2.5),
        ('LEFTPADDING', (0,0), (-1,-1), 4),
        ('RIGHTPADDING', (0,0), (-1,-1), 4),
        ('ALIGN', (0,0), (0,-1), 'CENTER'),
        ('ALIGN', (1,0), (1,-1), 'CENTER'),
        ('ALIGN', (3,0), (3,-1), 'CENTER'),
        ('ALIGN', (5,0), (5,-1), 'CENTER'),
        ('ALIGN', (6,0), (6,-1), 'CENTER'),
    ]
    for r in range(1, len(table_data)):
        if r % 2 == 0:
            ts.append(('BACKGROUND', (0, r), (-1, r), colors.HexColor('#F8FAFC')))
    t.setStyle(TableStyle(ts))
    story.append(t)
    story.append(Spacer(1, 5))

    # Instructions footer
    instructions = Paragraph(
        "<b>Important Instructions for Students:</b> 1. Report to your allotted exam room at least 15 minutes before the exam commences. "
        "2. Verify your roll number and name on the desk label before sitting. "
        "3. Strict alternating seating is maintained to ensure exam integrity. 4. Bring school ID and required stationery.",
        instruction_style
    )
    story.append(instructions)

    doc.build(story)
    bio.seek(0)
    pdf_bytes = bio.getvalue()
    pdf_filename = f"Exam_Dispatch_{sec_slug}.pdf"

    return pdf_bytes, pdf_filename, len(class_students), list(room_splits.values())

@router.get("/dispatch-pdf/{grade_class}")
def get_class_dispatch_pdf(
    grade_class: str,
    session_id: Optional[str] = Query(None, description="Optional exam session ID"),
    download: bool = Query(False, description="If true, returns attachment download, else inline display"),
    db: Session = Depends(get_db)
):
    """Generates and returns the clean Class Outward Exam Dispatch PDF for a section."""
    try:
        pdf_bytes, filename, student_count, room_splits = generate_class_dispatch_pdf(grade_class, session_id, db)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to generate dispatch PDF: {str(e)}")

    disposition = "attachment" if download else "inline"
    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={
            "Content-Disposition": f'{disposition}; filename="{filename}"',
            "Cache-Control": "no-cache"
        }
    )

@router.get("/whatsapp/notices")
def get_whatsapp_broadcast_notices(
    session_id: Optional[str] = Query(None, description="Optional exam session ID"),
    db: Session = Depends(get_db)
):
    """Returns compiled room-split data and ready-to-use WhatsApp notice texts for all class sections."""
    session = None
    if session_id:
        session = db.query(models.ExamSession).filter(models.ExamSession.id == session_id).first()
    if not session:
        alloc = db.query(models.SeatAllocation).first()
        if alloc:
            session = alloc.session
        else:
            session = db.query(models.ExamSession).first()

    session_name = session.name if session else "Exam Session"
    session_date = session.date if session else datetime.datetime.now().strftime("%d-%b-%Y")
    time_slot_str = session.time_slot if (session and session.time_slot) else "Morning (09:30 AM - 12:30 PM)"

    all_sections = [
        "XI - A", "XI - B", "XI - C", "XI - D", "XI - E", "XI - F", "XI - G", "XI - H",
        "XII - A", "XII - B", "XII - C", "XII - D", "XII - E", "XII - F", "XII - G"
    ]

    results = []
    for sec in all_sections:
        students = db.query(models.Student).filter(models.Student.grade == sec).order_by(models.Student.roll_no.asc()).all()
        stud_ids = [s.id for s in students]
        allocs = []
        if session and stud_ids:
            allocs = db.query(models.SeatAllocation).filter(
                models.SeatAllocation.session_id == session.id,
                models.SeatAllocation.student_id.in_(stud_ids)
            ).all()

        alloc_map = {a.student_id: a for a in allocs}

        # Room splits
        room_splits = {}
        subjects_set = set()
        for st in students:
            al = alloc_map.get(st.id)
            r_name = al.room.name if (al and al.room) else "Unassigned"
            if r_name not in room_splits:
                room_splits[r_name] = {"room": r_name, "count": 0, "rolls": []}
            room_splits[r_name]["count"] += 1
            room_splits[r_name]["rolls"].append(st.roll_no)
            if al and al.subject:
                subjects_set.add(al.subject.name)

        split_details = []
        for r_name, info in room_splits.items():
            numeric_rolls = [int(r) for r in info["rolls"] if str(r).isdigit()]
            if numeric_rolls:
                roll_range_str = f"Rolls {min(numeric_rolls)} to {max(numeric_rolls)}"
            else:
                roll_range_str = f"{info['count']} students"
            split_details.append({
                "room": r_name,
                "count": info["count"],
                "rollRange": roll_range_str
            })

        subject_display = ", ".join(sorted(subjects_set)) if subjects_set else "Scheduled Exam"

        # Format WhatsApp message text
        msg_lines = [
            "🏫 *CS ACADEMY - EXAM SEATING NOTICE*",
            "📋 *Class Outward Exam Despatch*",
            "━━━━━━━━━━━━━━━━━━━━━━━━━━",
            f"📅 *Session:* {session_name} ({session_date})",
            f"⏰ *Time:* {time_slot_str}",
            f"👥 *Class / Section:* {sec} ({len(students)} Students)",
            f"📚 *Subject:* {subject_display}",
            "━━━━━━━━━━━━━━━━━━━━━━━━━━",
            "📍 *EXAM ROOM ALLOCATIONS:*"
        ]

        for s in split_details:
            msg_lines.append(f"🔹 *Room {s['room']}* ({s['count']} students):")
            msg_lines.append(f"   • {s['rollRange']}")

        msg_lines.extend([
            "━━━━━━━━━━━━━━━━━━━━━━━━━━",
            "⚠️ *Important Instructions:*",
            "1. Report to your assigned room 15 minutes before exam time.",
            "2. Verify your roll number and name on your desk sticker.",
            "3. Strict 50/50 alternating seating is enforced across rooms.",
            "4. Bring your school ID card and required stationery.",
            "━━━━━━━━━━━━━━━━━━━━━━━━━━"
        ])

        whatsapp_text = "\n".join(msg_lines)
        encoded_sec = urllib.parse.quote(sec)
        sess_param = f"?session_id={session.id}" if session else ""
        pdf_url = f"/api/email/dispatch-pdf/{encoded_sec}{sess_param}"
        whatsapp_url = f"https://api.whatsapp.com/send?text={urllib.parse.quote(whatsapp_text)}"

        results.append({
            "section": sec,
            "gradeLevel": "Grade XI" if "XI" in sec and "XII" not in sec else "Grade XII",
            "studentCount": len(students),
            "allocatedCount": len(allocs),
            "subject": subject_display,
            "roomSplits": split_details,
            "pdfUrl": pdf_url,
            "whatsappText": whatsapp_text,
            "whatsappUrl": whatsapp_url
        })

    return {
        "status": "success",
        "sessionId": session.id if session else None,
        "sessionName": session_name,
        "sessionDate": session_date,
        "sessionTime": time_slot_str,
        "sections": results
    }
